# -*- coding: utf-8 -*-
"""数据访问层：候选池、隐藏结果、约束、元素先验"""
import csv
from pathlib import Path

import openpyxl

DATA = Path(r"D:\her_agent_project\data\Ni_RC_180_candidate_dataset_v2.xlsx")
CONSTRAINTS = Path(r"D:\her_agent_project\constraints.csv")
PRIORS = Path(r"D:\her_agent_project\element_priors.csv")

META = {"候选ID", "Ni组合方案", "组合元素"}


def load_pool():
    """只读'掩盖结果'：给 Agent 看的候选池，不含任何结果"""
    wb = openpyxl.load_workbook(DATA, data_only=True)
    ws = wb["掩盖结果"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    elems = [h for h in header if h and h not in META]

    pool = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(header, row))
        pool.append({
            "candidate_id": d["候选ID"],
            "composition": d["Ni组合方案"],
            "fractions": {e: float(d.get(e) or 0) for e in elems},
        })
    return pool


def load_results():
    """只给 execute 节点用：真实 Δη 与排名"""
    wb = openpyxl.load_workbook(DATA, data_only=True)
    ws = wb["完整数据"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    eta_col = next(h for h in header if h and "Δη" in h)

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(header, row))
        out[d["候选ID"]] = {
            "delta_eta_mv": float(d[eta_col]),
            "rank": int(d["全局排名"]),
            "top10": d["Top 10%"] == "是",
        }
    return out


def constraints_text():
    rows = list(csv.DictReader(open(CONSTRAINTS, encoding="utf-8-sig")))
    return "\n".join(
        f"{r['rule_id']} {r['rule_text']}（阈值 {r['threshold']}）"
        for r in rows
    )


def priors_text():
    rows = list(csv.DictReader(open(PRIORS, encoding="utf-8-sig")))
    return "\n".join(
        f"{r['symbol']}: 活性 {r['her_activity_prior']}，"
        f"碱稳定性 {r['alkaline_stability_prior']}，作用 {r['role']}"
        for r in rows
    )